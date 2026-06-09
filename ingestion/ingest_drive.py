#!/usr/bin/env python3
"""
GuidedBroker — full Resource Center ingestion (per-file links).

Reads ingestion/resource_center_manifest.csv (produced by DRIVE_MANIFEST.gs),
pulls each file's content straight from Drive, embeds it with
gemini-embedding-2-preview (768d, normalized), and upserts to Pinecone with that
file's OWN Drive share link attached as metadata. Each file also gets a "card"
vector (name + folder path) so it is findable by name even if its text can't be
extracted, and still returns the correct link.

Run:  python ingestion/ingest_drive.py
Env (from .env): GEMINI_API_KEY, PINECONE_API_KEY, PINECONE_INDEX_URL
Optional deps for Office files:  pip3 install python-docx openpyxl
"""
import os, csv, io, sys, time, hashlib, urllib.request
import numpy as np
from pinecone import Pinecone
from google import genai
from google.genai import types


def _load_env():
    p = os.path.join(os.path.dirname(__file__), "..", ".env")
    if os.path.exists(p):
        for line in open(p, encoding="utf-8"):
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, v = line.split("=", 1)
                os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))
_load_env()

GEMINI_API_KEY     = os.environ["GEMINI_API_KEY"]
PINECONE_API_KEY   = os.environ["PINECONE_API_KEY"]
PINECONE_INDEX_URL = os.environ["PINECONE_INDEX_URL"]
MANIFEST = os.path.join(os.path.dirname(__file__), "resource_center_manifest.csv")
DIM = 768

gem = genai.Client(api_key=GEMINI_API_KEY)
index = Pinecone(api_key=PINECONE_API_KEY).Index(host=PINECONE_INDEX_URL)


def normalize(v):
    a = np.array(v, dtype=np.float32); n = np.linalg.norm(a)
    return (a / n).tolist() if n > 0 else a.tolist()


import re

def embed(texts, task, batch=20):
    """Embed with automatic backoff on rate limits (429). Paced for free tier."""
    out = []
    n = len(texts)
    i = 0
    while i < n:
        chunk_texts = texts[i:i + batch]
        for attempt in range(10):
            try:
                r = gem.models.embed_content(
                    model="gemini-embedding-2-preview", contents=chunk_texts,
                    config=types.EmbedContentConfig(task_type=task, output_dimensionality=DIM))
                out.extend(normalize(e.values) for e in r.embeddings)
                break
            except Exception as ex:
                msg = str(ex)
                if "429" in msg or "RESOURCE_EXHAUSTED" in msg or "quota" in msg.lower():
                    m = re.search(r"retry in ([0-9.]+)s", msg) or re.search(r"([0-9.]+)s", msg)
                    delay = (float(m.group(1)) if m else 20) + 3
                    print("   rate limited — waiting %.0fs (%d/%d done)..." % (delay, i, n))
                    time.sleep(delay)
                    continue
                raise
        else:
            raise RuntimeError("embedding kept failing after 10 retries")
        i += batch
        print("   embedded %d/%d" % (min(i, n), n))
        time.sleep(1.0)  # gentle pacing so we don't slam the free-tier limit
    return out


def http_get(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=90) as resp:
        return resp.read()


def fetch_text(fid, mime, name):
    """Pull a file's text content from Drive using public (anyone-with-link) endpoints."""
    low = name.lower()
    try:
        if mime == "application/pdf" or low.endswith(".pdf"):
            data = http_get("https://drive.google.com/uc?export=download&id=" + fid)
            from pypdf import PdfReader
            r = PdfReader(io.BytesIO(data))
            return "\n\n".join((p.extract_text() or "") for p in r.pages)
        if mime == "application/vnd.google-apps.document":
            return http_get("https://docs.google.com/document/d/%s/export?format=txt" % fid).decode("utf-8", "ignore")
        if mime == "application/vnd.google-apps.spreadsheet":
            return http_get("https://docs.google.com/spreadsheets/d/%s/export?format=csv" % fid).decode("utf-8", "ignore")
        if mime == "application/vnd.google-apps.presentation":
            return http_get("https://docs.google.com/presentation/d/%s/export?format=txt" % fid).decode("utf-8", "ignore")
        if mime in ("text/plain", "text/csv") or low.endswith((".txt", ".md", ".csv")):
            return http_get("https://drive.google.com/uc?export=download&id=" + fid).decode("utf-8", "ignore")
        if low.endswith(".docx") or "wordprocessingml" in mime:
            data = http_get("https://drive.google.com/uc?export=download&id=" + fid)
            try:
                import docx
                return "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
            except Exception:
                return ""
        if low.endswith(".xlsx") or "spreadsheetml" in mime:
            data = http_get("https://drive.google.com/uc?export=download&id=" + fid)
            try:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
                lines = []
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        lines.append(",".join("" if c is None else str(c) for c in row))
                return "\n".join(lines)
            except Exception:
                return ""
    except Exception as e:
        print("   ! could not fetch %s: %s" % (name, e))
    return ""


def chunk(t, size=1800, overlap=200):
    out, i = [], 0
    while i < len(t):
        out.append(t[i:i + size]); i += size - overlap
    return [c for c in out if c.strip()]


def vid(i, *p):
    return hashlib.sha1((str(i) + "::" + "::".join(p)).encode()).hexdigest()[:24]


def main():
    if not os.path.exists(MANIFEST):
        print("Missing %s. Run DRIVE_MANIFEST.gs first and drop the CSV here." % MANIFEST)
        sys.exit(1)

    rows = list(csv.DictReader(open(MANIFEST, encoding="utf-8")))
    print("Manifest files: %d" % len(rows))

    items = []  # (kind, text, title, url, section, directions, file_name)
    for row in rows:
        name = row["name"]; fid = row["fileId"]; mime = row.get("mimeType", "")
        url = row.get("shareUrl", "#"); path = row.get("path", "")
        parts = path.split("/")
        section = parts[1] if len(parts) > 1 else path
        directions = 'In the Resource Center, open "%s" and select "%s".' % (path, name)

        # Always add a card vector (name + location) -> findable + correct link.
        items.append(("card", "%s. Located in %s." % (name, path), name, url, section, directions, name))

        text = fetch_text(fid, mime, name)
        n_chunks = 0
        if text.strip():
            for c in chunk(text):
                items.append(("doc", c, name, url, section, directions, name))
                n_chunks += 1
        print("  + %-55s %d chunks" % (name[:55], n_chunks))

    print("Embedding %d vectors (this is the slow part on the free tier)..." % len(items))
    embs = embed([it[1] for it in items], "RETRIEVAL_DOCUMENT")

    payload = []
    for i, ((kind, txt, title, url, section, directions, fname), e) in enumerate(zip(items, embs)):
        payload.append({"id": vid(i, kind, title), "values": e, "metadata": {
            "text": txt[:1500], "title": title, "url": url, "category": section,
            "directions": directions, "file_name": fname, "kind": kind}})

    # Only now that embedding succeeded do we clear the old vectors and load new.
    try:
        index.delete(delete_all=True)
        print("Cleared old vectors.")
        time.sleep(2)
    except Exception as e:
        print("Index clear skipped:", e)

    for i in range(0, len(payload), 100):
        index.upsert(vectors=payload[i:i + 100])
        print("Upserted %d/%d" % (min(i + 100, len(payload)), len(payload)))
    print("Done. %d vectors live in Pinecone." % len(payload))


if __name__ == "__main__":
    main()
